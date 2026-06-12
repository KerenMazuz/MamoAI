"""
Builds the financial planning output workbook from the organizer/advisor
agent results: one sheet per category, plus a balance overview and an
insights/recommendations report sheet.
"""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")
RTL_ALIGN = Alignment(horizontal="right", readingOrder=2, wrap_text=True)


def _write_table(ws, rows: list[dict], start_row: int = 1) -> int:
    """Write a list of dicts as a table starting at start_row. Returns next free row."""
    if not rows:
        ws.cell(row=start_row, column=1, value="(אין נתונים)").alignment = RTL_ALIGN
        return start_row + 2

    headers = list(rows[0].keys())
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = RTL_ALIGN

    for r, row in enumerate(rows, start=start_row + 1):
        for col, header in enumerate(headers, start=1):
            value = row.get(header)
            if isinstance(value, (list, dict)):
                value = ", ".join(str(v) for v in value) if isinstance(value, list) else str(value)
            cell = ws.cell(row=r, column=col, value=value)
            cell.alignment = RTL_ALIGN

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 22

    ws.sheet_view.rightToLeft = True
    return start_row + len(rows) + 2


def _write_title(ws, title: str, row: int) -> int:
    cell = ws.cell(row=row, column=1, value=title)
    cell.font = Font(bold=True, size=13)
    cell.alignment = RTL_ALIGN
    ws.sheet_view.rightToLeft = True
    return row + 1


def _write_kv(ws, data: dict, start_row: int = 1) -> int:
    row = start_row
    for key, value in data.items():
        ws.cell(row=row, column=1, value=key).alignment = RTL_ALIGN
        cell = ws.cell(row=row, column=2, value=value)
        cell.alignment = RTL_ALIGN
        row += 1
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 22
    ws.sheet_view.rightToLeft = True
    return row + 1


def build_workbook(organized: dict, advice: dict) -> bytes:
    """
    Args:
        organized: output of FinancialOrganizer.organize()
        advice: output of FinancialAdvisor.advise()

    Returns:
        bytes of the .xlsx file
    """
    wb = Workbook()

    # ── מאזן כללי ──────────────────────────────────────────────
    ws = wb.active
    ws.title = "מאזן כללי"
    row = _write_title(ws, "מאזן חודשי כולל", 1)
    row += 1
    summary = advice.get("balance_summary", {})
    row = _write_kv(ws, {
        "סך הכנסות חודשיות": summary.get("total_monthly_income"),
        "סך הוצאות חודשיות (כולל מחזוריות מחולקות ל-12)": summary.get("total_monthly_expenses"),
        "עודף/גירעון חודשי": summary.get("monthly_surplus_or_deficit"),
        "אחוז חיסכון": summary.get("savings_rate_percent"),
    }, row + 1)

    # ── הכנסות ─────────────────────────────────────────────────
    ws = wb.create_sheet("הכנסות")
    income = organized.get("income", {})
    row = _write_title(ws, "הכנסות שוטפות חודשיות", 1) + 1
    row = _write_table(ws, income.get("monthly_recurring", []), row)
    row = _write_title(ws, "הכנסות חד-פעמיות", row) + 1
    _write_table(ws, income.get("one_time", []), row)

    # ── הוצאות שוטפות ──────────────────────────────────────────
    expenses = organized.get("expenses", {})
    ws = wb.create_sheet("הוצאות שוטפות")
    row = _write_title(ws, "הוצאות שוטפות חודשיות", 1) + 1
    row = _write_table(ws, expenses.get("ongoing_monthly", []), row)

    credit_card = organized.get("credit_card_summary")
    if credit_card:
        row = _write_title(ws, "פירוט הוצאות כרטיס אשראי לפי קטגוריה", row) + 1
        row = _write_kv(ws, {"סך הוצאה חודשית בכרטיס": credit_card.get("total_monthly_spend")}, row)
        _write_table(ws, credit_card.get("by_category", []), row)

    # ── הוצאות מחזוריות ומגורים ────────────────────────────────
    ws = wb.create_sheet("הוצאות מחזוריות ומגורים")
    row = _write_title(ws, "הוצאות מחזוריות / מגורים", 1) + 1
    _write_table(ws, expenses.get("periodic_housing", []), row)

    # ── הוצאות עתידיות חד-פעמיות ───────────────────────────────
    ws = wb.create_sheet("הוצאות עתידיות")
    row = _write_title(ws, "הוצאות עתידיות חד-פעמיות / פעם בחיים", 1) + 1
    row = _write_table(ws, expenses.get("future_one_time", []), row)
    row = _write_title(ws, "תוכנית היערכות (מהיועץ)", row) + 1
    ws.cell(row=row, column=1, value=advice.get("future_expenses_plan", "")).alignment = RTL_ALIGN

    # ── חיסכון, קרן חירום ופנסיה ───────────────────────────────
    ws = wb.create_sheet("חיסכון, חירום ופנסיה")
    savings = organized.get("savings", {})
    row = _write_title(ws, "חיסכון שוטף", 1) + 1
    row = _write_kv(ws, {"חיסכון חודשי בתקציב": savings.get("monthly_savings_in_budget")}, row)
    row = _write_title(ws, "קרן חירום", row) + 1
    row = _write_kv(ws, savings.get("emergency_fund", {}) or {}, row)
    row = _write_title(ws, "פנסיה ופרישה", row) + 1
    row = _write_kv(ws, savings.get("pension_retirement", {}) or {}, row)
    row = _write_title(ws, "הערכת קרן חירום (מהיועץ)", row) + 1
    ws.cell(row=row, column=1, value=advice.get("emergency_fund_assessment", "")).alignment = RTL_ALIGN
    row += 2
    row = _write_title(ws, "הערכת פנסיה ופרישה (מהיועץ)", row) + 1
    ws.cell(row=row, column=1, value=advice.get("retirement_assessment", "")).alignment = RTL_ALIGN

    # ── ביטוחים ────────────────────────────────────────────────
    ws = wb.create_sheet("ביטוחים")
    insurance = organized.get("insurance", {})
    row = _write_title(ws, "ביטוחי בריאות וחיים שוטפים", 1) + 1
    row = _write_table(ws, insurance.get("health_and_life", []), row)
    row = _write_title(ws, "ביטוחי הגנת הון", row) + 1
    row = _write_table(ws, insurance.get("asset_protection", []), row)
    row = _write_title(ws, "פערים ביטוחיים (מהיועץ)", row) + 1
    for gap in advice.get("insurance_gaps", []):
        ws.cell(row=row, column=1, value=f"• {gap}").alignment = RTL_ALIGN
        row += 1

    # ── הון חופשי ──────────────────────────────────────────────
    ws = wb.create_sheet("הון חופשי")
    row = _write_title(ws, "הון חופשי", 1) + 1
    row = _write_table(ws, organized.get("free_capital", []), row)
    row = _write_title(ws, "המלצה (מהיועץ)", row) + 1
    ws.cell(row=row, column=1, value=advice.get("free_capital_recommendation", "")).alignment = RTL_ALIGN

    # ── קרן הלוואה עתידית לעצמי ────────────────────────────────
    ws = wb.create_sheet("קרן הלוואה עתידית לעצמי")
    strategy = advice.get("self_loan_fund_strategy", {}) or {}
    row = _write_title(ws, "אסטרטגיית קרן ההלוואה העתידית לעצמי", 1) + 1
    row = _write_kv(ws, {
        "זרימה חודשית נוכחית לקרן": strategy.get("current_monthly_flow"),
        "זרימה חודשית מומלצת לקרן": strategy.get("recommended_monthly_flow"),
    }, row)
    ws.cell(row=row, column=1, value=strategy.get("rationale", "")).alignment = RTL_ALIGN
    ws.column_dimensions["A"].width = 100

    # ── תובנות והמלצות ─────────────────────────────────────────
    ws = wb.create_sheet("תובנות והמלצות")
    row = _write_title(ws, "סיכום מנהלים", 1) + 1
    ws.cell(row=row, column=1, value=advice.get("executive_summary_markdown", "")).alignment = RTL_ALIGN
    ws.row_dimensions[row].height = 400
    row += 2

    recommendations = advice.get("recommendations", {}) or {}
    rec_sections = [
        ("immediate", "המלצות מיידיות"),
        ("medium_term", "המלצות לטווח בינוני"),
        ("strategic", "המלצות אסטרטגיות"),
    ]
    for key, title in rec_sections:
        row = _write_title(ws, title, row) + 1
        for item in recommendations.get(key, []):
            ws.cell(row=row, column=1, value=f"• {item}").alignment = RTL_ALIGN
            row += 1
        row += 1

    row = _write_title(ws, "הגדלת הכנסות מול חלוקת התקציב", row) + 1
    ws.cell(row=row, column=1, value=advice.get("income_vs_allocation_recommendation", "")).alignment = RTL_ALIGN
    row += 2

    missing = advice.get("missing_data_followup", [])
    if missing:
        row = _write_title(ws, "מידע חסר שכדאי להשלים", row) + 1
        for item in missing:
            ws.cell(row=row, column=1, value=f"• {item}").alignment = RTL_ALIGN
            row += 1

    ws.column_dimensions["A"].width = 100

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
